import socket
import tempfile
import os
from fpdf import FPDF
from docx import Document

MASTER_HOST = "localhost"
MASTER_PORT = 9000

def convert_file(file_data, conversion_type, output_path):
    text = file_data.decode('utf-8', errors='ignore')
    if conversion_type == "txt_to_pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for line in text.split('\n'):
            pdf.cell(200, 10, txt=line, ln=True)
        pdf.output(output_path)
    elif conversion_type == "txt_to_docx":
        doc = Document()
        doc.add_paragraph(text)
        doc.save(output_path)
    elif conversion_type == "txt_to_rtf":
        rtf_header = r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs24"
        rtf_footer = "}"
        text_esc = text.replace('\\', '\\\\').replace('{', '\\{').replace('}', '\\}')
        rtf_body = text_esc.replace('\n', '\\par ')
        rtf_content = rtf_header + rtf_body + rtf_footer
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(rtf_content)
    elif conversion_type == "txt_to_html":
        html = f"<html><body><pre>{text}</pre></body></html>"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html)
    else:
        raise ValueError(f"Unsupported conversion: {conversion_type}")

def main():
    sock = socket.socket()
    sock.connect((MASTER_HOST, MASTER_PORT))
    sock.send(b'W')   # identify as worker
    print("[Worker] Connected to master, waiting for jobs...")

    while True:
        # Read length prefix (4 bytes)
        raw_len = sock.recv(4)
        if not raw_len:
            print("[Worker] Master closed connection")
            break
        data_len = int.from_bytes(raw_len, 'big')
        print(f"[Worker] Expecting {data_len} bytes")

        payload = b''
        while len(payload) < data_len:
            chunk = sock.recv(min(4096, data_len - len(payload)))
            if not chunk:
                break
            payload += chunk

        if b"||" not in payload:
            err_msg = b"ERROR: Invalid job format"
            sock.send(len(err_msg).to_bytes(4, 'big'))
            sock.send(err_msg)
            continue

        conv_type, file_data = payload.split(b"||", 1)
        conv_type = conv_type.decode()
        print(f"[Worker] Converting {conv_type} ({len(file_data)} bytes)")

        ext = conv_type.split("_to_")[1]
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp_out:
            tmp_out_path = tmp_out.name

        try:
            convert_file(file_data, conv_type, tmp_out_path)
            with open(tmp_out_path, 'rb') as f:
                result_data = f.read()
            # Send result back with length prefix
            sock.send(len(result_data).to_bytes(4, 'big'))
            sock.sendall(result_data)
            print(f"[Worker] Sent {len(result_data)} bytes back")
        except Exception as e:
            err_msg = f"ERROR: {str(e)}".encode()
            sock.send(len(err_msg).to_bytes(4, 'big'))
            sock.send(err_msg)
            print(f"[Worker] Conversion error: {e}")
        finally:
            os.unlink(tmp_out_path)

    sock.close()

if __name__ == "__main__":
    main()