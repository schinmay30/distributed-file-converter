import os
from fpdf import FPDF
from docx import Document

def convert_txt_to_pdf(input_data, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    text = input_data.decode('utf-8', errors='ignore')
    for line in text.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)
    pdf.output(output_path)

def convert_txt_to_docx(input_data, output_path):
    doc = Document()
    text = input_data.decode('utf-8', errors='ignore')
    doc.add_paragraph(text)
    doc.save(output_path)

def convert_txt_to_rtf(input_data, output_path):
    rtf_header = r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Arial;}}\f0\fs24"
    rtf_footer = "}"
    text = input_data.decode('utf-8', errors='ignore')
    text = text.replace('\\', '\\\\').replace('{', '\\{').replace('}', '\\}')
    rtf_body = text.replace('\n', '\\par ')
    rtf_content = rtf_header + rtf_body + rtf_footer
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(rtf_content)

def convert_txt_to_html(input_data, output_path):
    text = input_data.decode('utf-8', errors='ignore')
    html = f"<html><body><pre>{text}</pre></body></html>"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

CONVERTERS = {
    "txt_to_pdf": convert_txt_to_pdf,
    "txt_to_docx": convert_txt_to_docx,
    "txt_to_rtf": convert_txt_to_rtf,
    "txt_to_html": convert_txt_to_html,
}

def convert_file(input_data, conversion_type, output_path):
    if conversion_type not in CONVERTERS:
        raise ValueError(f"Unsupported conversion: {conversion_type}")
    CONVERTERS[conversion_type](input_data, output_path)